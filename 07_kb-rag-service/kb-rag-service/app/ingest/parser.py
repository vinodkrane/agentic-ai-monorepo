import io
import os
import re


class UnsupportedFile(Exception):
    pass


def parse(filename, data):
    """Turn raw bytes into plain text. Returns a list of (page_no, text)."""
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        pages = _parse_pdf(data)
    elif ext == ".docx":
        pages = [(1, _parse_docx(data))]
    elif ext in (".txt", ".md"):
        pages = [(1, data.decode("utf-8", errors="ignore"))]
    else:
        raise UnsupportedFile(ext)

    pages = [(n, clean(t)) for n, t in pages]
    return [(n, t) for n, t in pages if t]


def _parse_pdf(data):
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        pages.append((i, page.extract_text() or ""))
    # scanned pdfs come back empty here. we don't OCR yet, the handler
    # logs it so we can see how often it happens (see docs/challenges...)
    return pages


def _parse_docx(data):
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    # tables were being dropped silently, pull cell text in row order
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def clean(text):
    text = text.replace("\x00", "")
    # pdf extraction leaves "exam-\nple" hyphen breaks
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
